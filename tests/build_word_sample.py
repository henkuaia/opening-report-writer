"""Build the synthetic T6 sample from the recorded T2 response, not an answer fixture."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
response = (ROOT / "session-check.md").read_text(encoding="utf-8")
body = response.split("### 正文\n", 1)[1].split("### 对应依据", 1)[0].strip()
paragraphs = body.split("\n\n")
doc = Document()
doc.core_properties.author = "opening-report-writer synthetic test"
doc.core_properties.last_modified_by = ""
doc.core_properties.title = "预约同步系统虚构验收小样"
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(2.5)
section.left_margin = section.right_margin = Cm(2.5)

for name, size, east_asia in (
    ("Normal", 12, "宋体"),
    ("Title", 18, "黑体"),
    ("Heading 1", 14, "黑体"),
):
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_after = Pt(6)
doc.styles["Normal"].paragraph_format.first_line_indent = Pt(24)

doc.add_paragraph("预约同步系统虚构验收小样", "Title")
doc.add_paragraph("本文件仅用于写作与排版验收。材料甲、乙、丙为虚构摘录，不是实际文献或系统实测结果。")
doc.add_heading("研究背景", 1)
for text in paragraphs:
    doc.add_paragraph(text)

caption = doc.add_paragraph("表1 三路线条件对照")
caption.paragraph_format.first_line_indent = Pt(0)
caption.paragraph_format.keep_with_next = True
rows = [
    ("材料", "断网可执行操作", "确认前提"),
    ("甲", "查询缓存，不能提交新预约", "在线访问中心并获确认"),
    ("乙", "登记待确认预约", "恢复后由中心判定，冲突可被拒绝"),
    ("丙", "仅在获授时段内离线确认", "预先获得互不重叠时段的授权"),
]
table = doc.add_table(rows=0, cols=3)
table.autofit = False
for column, width in zip(table.columns, (2, 7, 7)):
    column.width = Cm(width)
for row_number, values in enumerate(rows):
    cells = table.add_row().cells
    for cell, value, width in zip(cells, values, (2, 7, 7)):
        cell.width = Cm(width)
        cell.text = value
        cell.paragraphs[0].paragraph_format.first_line_indent = Pt(0)
        properties = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            border = OxmlElement(f"w:{side}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:color"), "D9D9D9")
            borders.append(border)
        properties.append(borders)
        if row_number == 0:
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "EEEEEE")
            properties.append(shade)
    if row_number == 0:
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

sections = [
    ("研究内容", "拟比较意向登记与预约确认两种操作语义，明确离线可用范围及重连后的状态转换。最终采用何种策略仍待真实需求与文献确定，不以该小样宣称已有系统成果。"),
    ("验证方案", "拟在相同预约操作序列下检查正常连接、断连恢复和重复提交时的状态变化，并检验冲突请求的反馈。网络条件、业务排他规则和性能阈值待补，不预填实验数据或训练流程。"),
    ("进度", "拟先确认业务规则与弱网条件，再实现原型、开展功能和故障场景验证，最后整理分析与报告。截止日期尚未给出，暂不填写日历时间或已完成状态。"),
    ("参考资料", "材料甲第1段、材料乙第1段、材料丙第1段均为本次虚构验收输入；材料编号用于追溯摘录，不是正式参考文献。真实报告需另补相应来源及完整题录。"),
]
for heading, text in sections:
    doc.add_heading(heading, 1)
    doc.add_paragraph(text)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.first_line_indent = Pt(0)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
target = ROOT / "word-sample.docx"
doc.save(target)

# These checks validate content and OOXML, not pagination or visual quality.
saved = Document(target)
assert [p.text for p in saved.paragraphs if p.style.name == "Heading 1"] == [
    "研究背景", "研究内容", "验证方案", "进度", "参考资料"
]
assert all(text in [p.text for p in saved.paragraphs] for text in paragraphs)
assert len(saved.tables) == 1 and len(saved.tables[0].rows) == 4
assert saved.styles["Normal"].font.size == Pt(12)
assert saved.styles["Normal"].paragraph_format.line_spacing == 1.5
assert "PAGE" in saved.sections[0].footer._element.xml
assert abs(saved.sections[0].left_margin.cm - 2.5) < 0.01
print("PASS: generated and parsed; T2 text, headings, table, 12pt, 1.5 spacing, margin and PAGE field")
print("NOT VERIFIED: rendering, visual layout, Microsoft Word opening")
